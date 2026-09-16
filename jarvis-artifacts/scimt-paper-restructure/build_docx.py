"""Build a .docx of the send-out doc from template.html, for upload to Google Drive.

Includes sections 1-4, Appendix A, and a written tl;dr of Appendix B (passed in
as APPENDIX_B_TLDR below). Sidenotes become indented italic notes under their
paragraph; inline SVG figures are rasterised with cairosvg; PNG figures come from
the same lookup build.py uses. Usage::

    python3 build_docx.py --out stress_testing_midtraining_draft.docx
"""
from __future__ import annotations

import argparse, base64, io, pathlib, re, sys, tempfile
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import cairosvg

sys.path.insert(0, str(pathlib.Path(__file__).parent))
import build as B  # noqa: E402  (FIGS + DEFAULT_SCIMT)

HERE = pathlib.Path(__file__).parent

APPENDIX_B_TLDR = [
    ("Scaling midtraining dose and EFT dose.",
     "We midtrained Gemma 3 12B and 27B at 1M to 190M presented Charter tokens and elicited with one or two epochs of EFT. "
     "More midtraining helps monotonically and had not saturated at 190M tokens (Gemma 27B: 48% Charter at 5M, 75% at 190M, control 26 to 43%). "
     "One versus two EFT epochs is within the seed spread. A tenfold EFT scale-up on GLM-4.5-Air (81,920 episodes) keeps the held-in effect "
     "(81% versus 23% for the control) but does not restore the held-out clauses (22%)."),
    ("Worked examples in the midtraining corpus.",
     "Each clause appears in the corpus in a worked variant and a qualitative variant. Restricting the corpus to the qualitative documents at the same dose "
     "cuts the Charter arm from 65% to 37% after agreement-only EFT (Gemma 3 12B, 50M). Under 2% coin-labelled EFT both corpora collapse to single digits."),
    ("How much conflicting EFT is enough.",
     "We ran a ladder of conflict fractions from 0.5% to 5%. On Gemma 3 at every dose, 0.5% coin-labelled rows (41 of 8,192) already bring the Charter arm to 12 to 32%, and 5% to 2 to 6%. "
     "In the other direction, 1% Charter-labelled rows lift the control with no midtraining to 59 to 76%. The label on the conflict rows sets the endpoint; the prior only sets where the line starts."),
    ("SFT versus GRPO as the elicitation method.",
     "On Gemma-4-26B-A4B grafts, SFT on agreement episodes separates the Charter- and Coin-midtrained models by 32 points on conflict episodes; GRPO on the same episodes separates them by 10, "
     "against 12 before any post-training. The Charter-midtrained model still reasons about the Charter in its thinking traces before choosing the cheaper crew: the content is retrievable, RL does not make it load-bearing."),
    ("A second setting, Python 4.",
     "A fictional dialect of Python with eight syntax rules, four demonstrated by EFT and four held out. Here the held-out rules do transfer (40 to 85% adoption versus 0 to 16% for the control on Gemma 3 27B), "
     "though EFT on the other rules suppresses three of the four relative to the pre-EFT model. Our reading: midtraining installs independently retrievable habits, not a structured principle."),
    ("Reproducing Model Spec Midtraining.",
     "We reproduced the published recipe on six base models. It raises the targeted preference on most substrates (America: 35 to 60% with midtraining versus 17 to 32% without), but much of the lift is present without the alignment fine-tuning it is paired with, "
     "and the midtraining-by-fine-tuning interaction is significant on six of twelve arms and small where it is. The single-value setting cannot ask the held-out or ambiguity questions above, which is why we built Dispatch."),
]

# ------------------------------------------------------------------ helpers
def load_img(key: str, scimt: pathlib.Path) -> bytes:
    if key in B.FIGS:
        p = scimt / B.FIGS[key]
    else:
        cands = [scimt / "figures" / key.rsplit("_v", 1)[0] / f"{key}.png", scimt / "tex/Images" / f"{key}.png",
                 scimt / "tex/Images" / f"{key}.jpeg", *(scimt / "figures").glob(f"*/{key}.png")]
        p = next((c for c in cands if c.exists()), None)
    if p is None or not p.exists():
        raise FileNotFoundError(key)
    return p.read_bytes()


RAW_SVGS: list[str] = []


def raw_svg_for(label: str) -> str:
    """html.parser lowercases attribute names (viewBox, markerWidth...), which breaks cairosvg;
    take the SVG source verbatim from the template instead, matched by aria-label."""
    for raw in RAW_SVGS:
        if f'aria-label="{label}"' in raw:
            return raw
    raise KeyError(label)


def svg_png(svg: str) -> bytes:
    svg = svg.replace("currentColor", "#1a1d1b")
    m = re.search(r'viewBox="0 0 (\d+) (\d+)"', svg)
    if m and ' width=' not in svg[:300]:
        svg = svg.replace('<svg ', f'<svg width="{m.group(1)}" height="{m.group(2)}" ', 1)
    return cairosvg.svg2png(bytestring=svg.encode(), output_width=1800, background_color="white")


def add_runs(par, node, doc):
    """Append inline content of node to paragraph par; collect sidenotes."""
    notes = []
    for child in node.children:
        if isinstance(child, NavigableString):
            par.add_run(str(child))
        elif isinstance(child, Tag):
            cls = child.get("class", [])
            if "sidenote" in cls:
                notes.append(child.get_text(" ", strip=True))
            elif "sn-ref" in cls:
                r = par.add_run(f"[{len(notes) + 1}]"); r.font.superscript = True
            elif child.name in ("b", "strong"):
                r = par.add_run(child.get_text()); r.bold = True
            elif child.name in ("em", "i"):
                r = par.add_run(child.get_text()); r.italic = True
            elif child.name == "code":
                r = par.add_run(child.get_text()); r.font.name = "Courier New"
            elif child.name == "br":
                par.add_run("\n")
            else:
                notes += add_runs(par, child, doc)
    return notes


def add_notes(doc, notes):
    for i, n in enumerate(notes, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        r = p.add_run(f"[{i}] {n}"); r.italic = True; r.font.size = Pt(9)


def add_paragraph(doc, node, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    notes = add_runs(p, node, doc)
    add_notes(doc, notes)


def add_figure(doc, fig, caption, scimt):
    img = fig.find("img"); svg = fig.find("svg"); ph = fig.find(class_="placeholder")
    data = None
    if img is not None:
        m = re.match(r"\{\{IMG:([a-z0-9_]+)\}\}", img.get("src", ""))
        if m: data = load_img(m.group(1), scimt)
    elif svg is not None:
        data = svg_png(raw_svg_for(svg.get("aria-label", "")))
    if data:
        doc.add_picture(io.BytesIO(data), width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ph is not None:
        p = doc.add_paragraph(); r = p.add_run("[Figure placeholder]"); r.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption is not None:
        p = doc.add_paragraph()
        for child in caption.children:
            if isinstance(child, Tag) and "fig" in child.get("class", []):
                r = p.add_run(child.get_text() + ". "); r.bold = True
            elif isinstance(child, Tag) and child.name == "b":
                r = p.add_run(child.get_text()); r.bold = True
            elif isinstance(child, Tag) and "note" in child.get("class", []):
                r = p.add_run(child.get_text()); r.italic = True
            else:
                p.add_run(child.get_text() if isinstance(child, Tag) else str(child))
        for r in p.runs: r.font.size = Pt(9.5)


def add_table(doc, tbl):
    rows = tbl.find_all("tr")
    ncol = max(len(r.find_all(["td", "th"])) for r in rows)
    t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
    for r in rows:
        cells = r.find_all(["td", "th"]); row = t.add_row().cells
        for i, c in enumerate(cells):
            row[i].text = c.get_text(" ", strip=True)
            for par in row[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)
                    if c.name == "th": run.bold = True
    doc.add_paragraph()


def walk_section(doc, section, scimt, *, skip_ids=()):
    for node in section.children:
        if not isinstance(node, Tag): continue
        if node.get("id") in skip_ids: continue
        if node.name == "h2":
            doc.add_heading(node.get_text(" ", strip=True), level=1)
        elif node.name == "h3":
            doc.add_heading(node.get_text(" ", strip=True), level=2)
        elif node.name == "p":
            if "lede" in node.get("class", []): continue
            add_paragraph(doc, node)
        elif node.name in ("ul", "ol"):
            for li in node.find_all("li", recursive=False):
                add_paragraph(doc, li, style="List Bullet" if node.name == "ul" else "List Number")
        elif node.name == "figure":
            if "table" in node.get("class", []):
                cap = node.find("figcaption")
                if cap is not None: add_figure(doc, node, cap, scimt)  # caption only (no img)
                add_table(doc, node.find("table"))
            else:
                # caption is the next sibling figcaption
                sib = node.find_next_sibling()
                cap = sib if (sib is not None and sib.name == "figcaption") else None
                add_figure(doc, node, cap, scimt)
        elif node.name == "figcaption":
            continue  # consumed with its figure
        elif node.name == "blockquote":
            p = doc.add_paragraph(node.get_text(" ", strip=True)); p.paragraph_format.left_indent = Inches(0.4)
            for r in p.runs: r.italic = True
        elif node.name == "details":
            summ = node.find("summary"); pre = node.find("pre")
            if summ is not None:
                p = doc.add_paragraph(); r = p.add_run(summ.get_text(" ", strip=True)); r.bold = True
            if pre is not None:
                for line in pre.get_text().split("\n"):
                    p = doc.add_paragraph(line); p.paragraph_format.left_indent = Inches(0.3)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs: r.font.name = "Courier New"; r.font.size = Pt(8.5)
                doc.add_paragraph()
        elif node.name in ("article", "section", "div"):
            walk_section(doc, node, scimt, skip_ids=skip_ids)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scimt", type=pathlib.Path, default=B.DEFAULT_SCIMT)
    ap.add_argument("--out", type=pathlib.Path, default=HERE / "stress_testing_midtraining_draft.docx")
    a = ap.parse_args()
    html = (HERE / "template.html").read_text().split("{{PAPER}}")[0]
    RAW_SVGS.extend(re.findall(r"<svg .*?</svg>", html, re.S))
    soup = BeautifulSoup(html, "html.parser")
    main_el = soup.find("main")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Georgia"; st.font.size = Pt(10.5)
    doc.add_heading("Stress-Testing Alignment Midtraining", level=0)
    k = main_el.find(class_="kicker")
    if k is not None:
        p = doc.add_paragraph(k.get_text(" ", strip=True)); [setattr(r.font, "size", Pt(9)) for r in p.runs]
    # 1 Introduction (+ questions panel), hero figure
    walk_section(doc, main_el.find(class_="intro"), a.scimt)
    q = main_el.find(class_="questions")
    if q is not None:
        doc.add_heading(q.find("h2").get_text(" ", strip=True), level=2)
        for li in q.find_all("li"): add_paragraph(doc, li, style="List Number")
    hero = main_el.find("figure", class_="hero")
    if hero is not None:
        add_figure(doc, hero, hero.find_next_sibling("figcaption"), a.scimt)
    # sections 2, 3, 4, A
    for sid in ("setup", "results", "discussion", "appendix"):
        sec = main_el.find("section", id=sid)
        if sec is not None: walk_section(doc, sec, a.scimt)
    # Appendix B tl;dr
    doc.add_heading("B Additional analyses (summary)", level=1)
    doc.add_paragraph("We ran six further analyses. Full figures and numbers are in the web version of this document; "
                      "here is what each one asked and what it found.")
    for title, body in APPENDIX_B_TLDR:
        p = doc.add_paragraph(style="List Bullet"); r = p.add_run(title + " "); r.bold = True; p.add_run(body)
    doc.add_paragraph("Validity evaluations (capability, instruction following, preference coherence, refusal, perplexity) are in Appendix C of the web version; "
                      "midtraining changes little that a capability eval would notice.")
    doc.save(a.out)
    print("wrote", a.out, f"{a.out.stat().st_size/1e6:.1f} MB")


if __name__ == "__main__":
    main()
